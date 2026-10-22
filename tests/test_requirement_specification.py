import json
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

from sqlalchemy import create_engine, inspect
from app import create_app, db
from app.models import Project, Interview, Requirement
from app.schema_migrations import upgrade_requirement_specification
from app.services.requirement_specification import (
    extract_specification, validated_fields, specification_rows, relevant_context, PENDING,
)
from app.agents.requirements_agent import RequirementsAgent
from app.agents.srs_agent import SRSAgent
from app.services.document_service import DocumentService
from docx import Document


SOURCE = 'El bibliotecario debe registrar libros con título y autor. El título es obligatorio.'


class EvidenceTests(unittest.TestCase):
    def test_everyday_requests_and_unresolved_decisions(self):
        from app.services.functional_context import request, pending_information, classify_evidence
        from app.services.requirement_specification import FIELDS
        for sentence in ('También quisiera anotar devoluciones.', 'Me gustaría tener un resumen.',
                         'Lo más importante es anotar las ventas.'):
            self.assertTrue(request(sentence), sentence)
        for sentence in ('No he decidido cuántos días tendrá el cliente.',
                         'Lo del período no lo tengo claro.', 'Todavía no sé quién lo autorizará.'):
            self.assertTrue(pending_information(sentence))
            self.assertFalse(any(classify_evidence(sentence, FIELDS).values()))

    def test_everyday_rules_and_examples_keep_their_meaning(self):
        from app.services.functional_context import classify_evidence, request
        from app.services.requirement_specification import FIELDS
        source = ('Si quedan dos pantalones y alguien intenta vender tres, que avise que no alcanza la cantidad. '
                  'Por ejemplo, si tengo cinco camisetas y vendo dos, deberían quedar tres. '
                  'Buscar algo no debe cambiar las cantidades.')
        fields = classify_evidence(source, FIELDS)
        self.assertEqual(len(fields['alternatives']), 1)
        self.assertEqual(len(fields['acceptance_criteria']), 1)
        self.assertEqual(len(fields['business_rules']), 1)
        self.assertFalse(request('Buscar algo no debe cambiar las cantidades.'))

    def test_only_explicit_first_person_participation_identifies_actor(self):
        from app.services.functional_context import explicit_actor
        self.assertEqual(explicit_actor('Mi ayudante y yo anotamos las ventas.', ''),
                         'Persona entrevistada y su ayudante')
        self.assertEqual(explicit_actor('Yo sería quien lo consulte.', ''), 'Persona entrevistada')
        self.assertIsNone(explicit_actor('Quiero que se guarde una venta.', ''))

    def test_overview_keeps_goals_without_later_details(self):
        from app.services.functional_context import functional_blocks
        source = ('Necesitamos un sistema que permita registrar pacientes, agendar citas y emitir facturas. '
                  'El sistema deberá registrar pacientes con su nombre. '
                  'El sistema deberá agendar citas con fecha y hora.')
        blocks = functional_blocks(source)
        self.assertEqual(len(blocks), 3)
        self.assertTrue(any('facturas' in b['anchor'] for b in blocks))

    def test_undefined_criteria_do_not_turn_meetings_into_acceptance(self):
        from app.services.functional_context import classify_evidence
        from app.services.requirement_specification import FIELDS, coherent_fields
        source = 'Tampoco hemos acordado los criterios de aceptación. Necesitamos una reunión adicional para precisar esos aspectos.'
        self.assertFalse(any(classify_evidence(source, FIELDS).values()))
        self.assertFalse(any(coherent_fields({k: [source] for k in FIELDS}, source).values()))

    def test_multiple_goals_and_dependent_steps(self):
        from app.services.functional_context import functional_blocks
        blocks = functional_blocks('El sistema deberá registrar pacientes, agendar citas y emitir facturas.')
        self.assertEqual(len(blocks), 3)
        self.assertIn('registrar pacientes', blocks[0]['anchor'])
        self.assertEqual(blocks[1]['description'], 'El sistema deberá agendar citas.')
        self.assertEqual(len(functional_blocks('El sistema deberá registrar pacientes y mostrar un mensaje de confirmación.')), 1)

    def test_similar_words_do_not_discard_distinct_requirements(self):
        agent = RequirementsAgent()
        first = 'El sistema deberá registrar pagos de clientes.'
        self.assertFalse(agent._duplicate('El sistema deberá consultar pagos de clientes.', [{'description': first}]))
        self.assertFalse(agent._duplicate('El sistema deberá registrar pagos de proveedores.', [{'description': first}]))
        self.assertTrue(agent._duplicate(first, [{'description': first}]))

    def test_discovery_covers_all_batches_and_validates_indices(self):
        from app.services.functional_context import discover_starts
        llm = Mock()
        llm.generate_json.side_effect = [{'starts': [0, True, 9999, '1']}, {'starts': [1]}]
        self.assertEqual(discover_starts(['a' * 4000, 'b' * 4000], llm), {0, 1})
        self.assertEqual(llm.generate_json.call_count, 2)

    def test_development_fiche_keeps_components_separate(self):
        source = ('Antes de registrar el libro, el bibliotecario debe estar autenticado. '
                  'El bibliotecario debe registrar libros con título y autor. '
                  'Si el título está duplicado, se muestra un error.')
        req = {'description': 'Registrar libros', 'source_excerpt': source}
        llm = Mock()
        llm.generate_json.return_value = {'preconditions': [1], 'inputs': [2],
            'main_flow': [2], 'alternatives': [3], 'business_rules': [3]}
        spec = extract_specification(llm, req, source)
        rows = {r['key']: r for r in specification_rows(spec, source, req)}
        self.assertIn('autenticado', rows['preconditions']['values'][0])
        self.assertIn('título y autor', rows['inputs']['values'][0])
        self.assertIn('error', rows['alternatives']['values'][0])
        self.assertTrue(rows['main_flow']['values'])
        self.assertFalse(rows['acceptance_criteria']['values'])
        self.assertIn('observables', rows['acceptance_criteria']['help'])

    def test_library_details_stay_with_their_requirement(self):
        anchor = 'Los estudiantes deben poder consultar los libros disponibles por título o autor.'
        others = ['También queremos generar un reporte con los libros más prestados.',
                  'Sería bueno enviar una notificación cuando un préstamo esté próximo a vencer.',
                  'No necesitamos que el sistema permita comprar libros en línea.']
        source = ' '.join([anchor] + others)
        req = {'description': 'El sistema deberá consultar libros por título o autor.', 'source_excerpt': anchor}
        payload = {'fields': {key: [anchor] + others for key in
                              ('inputs', 'business_rules', 'outputs', 'acceptance_criteria')}}
        for row in specification_rows(payload, source, req):
            self.assertFalse(set(row['values']).intersection(others))
        llm = Mock()
        llm.generate_json.return_value = {'inputs': [1, 2, 3, 4]}
        fields = extract_specification(llm, req, source)['fields']
        self.assertEqual(fields['inputs'], [anchor])
        self.assertNotIn('comprar', llm.generate_json.call_args.args[0])

    def test_shared_action_does_not_link_separate_requests(self):
        anchor = 'El bibliotecario debe registrar libros con título y autor.'
        unrelated = 'También necesitamos registrar los préstamos de libros.'
        req = {'description': 'Registrar libros', 'source_excerpt': anchor}
        self.assertNotIn(unrelated, relevant_context(req, anchor + ' ' + unrelated))

    def test_function_statement_is_not_automatically_a_rule_or_result(self):
        anchor = 'Los estudiantes deben consultar libros por título o autor.'
        req = {'description': 'Consultar libros', 'source_excerpt': anchor}
        payload = {'fields': {key: [anchor] for key in
                              ('inputs', 'business_rules', 'outputs', 'acceptance_criteria')}}
        rows = specification_rows(payload, anchor, req)
        self.assertEqual(next(r['values'] for r in rows if r['key'] == 'inputs'), [anchor])
        self.assertFalse(next(r['values'] for r in rows if r['key'] == 'outputs'))
        self.assertFalse(next(r['values'] for r in rows if r['key'] == 'acceptance_criteria'))

    def test_coordinated_actions_are_not_assigned_together(self):
        anchor = 'El bibliotecario debe registrar nuevos libros y actualizar su información.'
        for description, expected in [('Registrar libros', 'registrar nuevos libros'),
                                      ('Actualizar información de libros', 'actualizar su información')]:
            scope = relevant_context({'description': description, 'source_excerpt': anchor}, anchor)
            self.assertIn(expected, scope)
            self.assertNotIn(anchor, scope)

    def test_rejected_feature_is_not_a_requirement_or_detail(self):
        source = 'No necesitamos que el sistema permita comprar libros en línea.'
        description = 'El sistema deberá no permitir la compra de libros en línea.'
        self.assertTrue(RequirementsAgent()._rejected(description, source))
        self.assertFalse(relevant_context({'description': description, 'source_excerpt': source}, source))
        self.assertFalse(RequirementsAgent()._rejected('Registrar libros', 'No registrar libros sin título.'))

    def test_drops_invented_rules_and_keeps_literal_evidence(self):
        fields = validated_fields({
            'inputs': ['título y autor', 'correo y contraseña'],
            'business_rules': ['El título es obligatorio.', 'El título debe tener entre 10 y 500 caracteres.'],
        }, SOURCE)
        self.assertEqual(fields['inputs'], ['título y autor'])
        self.assertEqual(fields['business_rules'], ['El título es obligatorio.'])
        self.assertNotIn('permissions', fields)

    def test_legacy_and_corrupt_payloads_are_pending(self):
        for payload in (None, '', '{bad', '[]', '{"fields": {"inputs": "inventado"}}'):
            self.assertTrue(all(not row['values'] for row in specification_rows(payload, SOURCE)))

    def test_compact_rows_merge_legacy_conditions_without_duplicates(self):
        spec = {'schema_version': 1, 'fields': {
            'inputs': ['El título es obligatorio.'],
            'business_rules': ['El título es obligatorio.'],
            'permissions': ['Solo el bibliotecario puede registrar.'],
            'postconditions': ['El libro queda registrado.'],
            'objective': ['Una finalidad que ya no se muestra.']}}
        original = json.dumps(spec)
        source = SOURCE + ' Solo el bibliotecario puede registrar. El libro queda registrado.'
        rows = specification_rows(spec, source)
        self.assertEqual(len(rows), 7)
        self.assertEqual(next(r['values'] for r in rows if r['key'] == 'business_rules'), ['El título es obligatorio.', 'Solo el bibliotecario puede registrar.'])
        self.assertEqual(next(r['values'] for r in rows if r['key'] == 'outputs'), ['El libro queda registrado.'])
        self.assertEqual(json.dumps(spec), original)

    def test_failure_keeps_pending_details(self):
        llm = Mock()
        llm.generate_json.side_effect = RuntimeError('offline')
        result = extract_specification(llm, {'description': 'Registrar libros', 'source_excerpt': SOURCE}, SOURCE)
        self.assertEqual(result['extraction_status'], 'unavailable')
        self.assertTrue(result['fields']['inputs'])
        self.assertFalse(result['fields']['outputs'])

    def test_details_have_no_academic_input_and_no_invented_approval(self):
        llm = Mock()
        llm.generate_json.return_value = {'inputs': [1, 99, True, 'inventado'], 'approved_by': 'Director'}
        result = extract_specification(llm, {'description': 'Registrar libros', 'source_excerpt': SOURCE}, SOURCE)
        self.assertNotIn('approved_by', result['fields'])
        self.assertEqual(result['schema_version'], 3)
        self.assertEqual(set(llm.generate_json.call_args.args[1]['properties']), {'preconditions', 'inputs', 'main_flow', 'business_rules', 'alternatives', 'outputs', 'acceptance_criteria'})
        self.assertEqual(result['fields']['inputs'], ['El bibliotecario debe registrar libros con título y autor.', 'El título es obligatorio.'])

    def test_migration_preserves_existing_rows_and_is_idempotent(self):
        engine = create_engine('sqlite://')
        with engine.begin() as conn:
            conn.exec_driver_sql('CREATE TABLE requirement (id INTEGER PRIMARY KEY, description TEXT)')
            conn.exec_driver_sql("INSERT INTO requirement VALUES (1, 'RF anterior')")
        upgrade_requirement_specification(engine)
        upgrade_requirement_specification(engine)
        with engine.connect() as conn:
            self.assertEqual(conn.exec_driver_sql('SELECT description, specification FROM requirement').one(), ('RF anterior', None))
        self.assertEqual(sum(c['name'] == 'specification' for c in inspect(engine).get_columns('requirement')), 1)
        engine.dispose()


class IntegrationTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory()
        root = Path(self.temp.name)
        self.app = create_app({'TESTING': True, 'SQLALCHEMY_DATABASE_URI': 'sqlite:///' + str(root/'test.db'),
                               'UPLOAD_FOLDER': str(root/'uploads'), 'GENERATED_FOLDER': str(root/'generated')})
        self.context = self.app.app_context()
        self.context.push()
        project = Project(name='Biblioteca', description='Préstamo de libros')
        db.session.add(project)
        db.session.flush()
        self.project_id = project.id
        interview = Interview(project_id=project.id, transcription=SOURCE)
        db.session.add(interview)
        db.session.commit()
        self.interview_id = interview.id

    def tearDown(self):
        db.session.remove()
        db.engine.dispose()
        self.context.pop()
        self.temp.cleanup()

    def save_requirement(self, specification=None):
        req = Requirement(project_id=self.project_id, interview_id=self.interview_id,
                          code='RF-001', description='El sistema deberá registrar libros con título y autor.',
                          source_excerpt=SOURCE, specification=specification)
        db.session.add(req)
        db.session.commit()
        return req

    def test_web_renders_legacy_and_escapes_evidence(self):
        req = self.save_requirement()
        response = self.app.test_client().get(f'/interviews/{self.interview_id}/requirements')
        self.assertEqual(response.status_code, 200)
        self.assertIn('Pendiente de aclaración', response.get_data(as_text=True))
        self.assertIn('antes de incorporar', response.get_data(as_text=True))
        interview = db.session.get(Interview, self.interview_id)
        malicious = 'El bibliotecario debe registrar libros con título <script>alert(1)</script>.'
        interview.transcription = malicious
        req.source_excerpt = malicious
        req.specification = json.dumps({'fields': {'inputs': [malicious]}})
        db.session.commit()
        html = self.app.test_client().get(f'/interviews/{self.interview_id}/requirements').get_data(as_text=True)
        self.assertIn('&lt;script&gt;', html)
        self.assertNotIn('<script>alert(1)</script>', html)

    def test_compact_ui_retains_controls_and_collapses_review(self):
        self.save_requirement(json.dumps({'schema_version': 1, 'fields': {'inputs': ['título y autor']}}))
        html = self.app.test_client().get(f'/interviews/{self.interview_id}/requirements').get_data(as_text=True)
        for label in ('Datos de entrada y validaciones', 'Resultado esperado', 'Criterios de aceptación', 'Analizar nuevamente', 'Generar SRS IEEE 830'):
            self.assertIn(label, html)
        self.assertIn('<details class=', html)
        self.assertNotIn('<details open', html)
        self.assertIn('id="project-context"', html)
        for label in ('Objetivo y necesidad', 'Flujo principal', 'Datos de gestión', 'Fecha de aprobación', 'Responsable de validación'):
            self.assertNotIn(label, html)

    def test_existing_unrelated_details_are_filtered_in_web_and_srs(self):
        unrelated = 'También queremos generar un reporte con los libros más prestados.'
        req = self.save_requirement(json.dumps({'fields': {
            'inputs': ['título y autor', unrelated], 'business_rules': [unrelated],
            'outputs': [unrelated], 'acceptance_criteria': [unrelated]}}))
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription += ' ' + unrelated
        db.session.commit()
        html = self.app.test_client().get(f'/interviews/{self.interview_id}/requirements').get_data(as_text=True)
        self.assertNotIn(unrelated, html)
        self.assertIn('título y autor', html)
        item = {'code': req.code, 'description': req.description, 'source_excerpt': req.source_excerpt,
                'specification': req.specification, 'interview_id': self.interview_id}
        project = {'name': 'Biblioteca', 'interviews': [{'id': self.interview_id, 'transcription': interview.transcription}]}
        content = SRSAgent()._build_document(project, [item], {})
        self.assertNotIn(unrelated, content)
        self.assertIn('título y autor', content)


    def test_srs_and_word_include_details_without_new_rules(self):
        spec = {'fields': {'inputs': ['título y autor'], 'business_rules': ['El título es obligatorio.'],
                           'acceptance_criteria': ['El título es obligatorio.'], 'permissions': ['Solo administradores.']}}
        item = {'code': 'RF-001', 'description': 'El sistema deberá registrar libros.',
                'interview_id': self.interview_id, 'source_excerpt': SOURCE, 'specification': json.dumps(spec)}
        project = {'name': 'Biblioteca', 'description': 'Libros', 'interviews': [{'id': self.interview_id, 'transcription': SOURCE}]}
        content = SRSAgent()._build_document(project, [item], {})
        self.assertIn('IEEE 830-1998', content)
        self.assertIn('**Datos de entrada y validaciones:**\n\n- título y autor', content)
        self.assertIn('RF-001-CA-01: El título es obligatorio.', content)
        self.assertNotIn('Solo administradores.', content)
        self.assertNotIn('**Responsable de validación:**', content)
        self.assertNotIn('**Fecha de aprobación:**', content)
        self.assertIn('**Resultado esperado:**', content)
        doc = Document(DocumentService().create_docx(content))
        self.assertTrue(any('RF-001-CA-01' in p.text for p in doc.paragraphs))
        self.assertTrue(any('título y autor' in p.text for p in doc.paragraphs))

    def test_mcp_transfers_new_and_legacy_specification(self):
        from app.mcp.servers import database_server
        req = self.save_requirement(json.dumps({'fields': {'inputs': ['título y autor']}}))
        with patch.object(database_server, 'DATABASE_PATH', db.engine.url.database):
            fetched = database_server.get_requirements(self.interview_id)['requirements'][0]
            self.assertEqual(fetched['specification'], req.specification)
            project = database_server.get_project_srs_data(self.project_id)['project']
            content = SRSAgent()._build_document(project, project['requirements'], {})
            self.assertIn('- título y autor', content)
            req.specification = None
            db.session.commit()
            fetched = database_server.get_requirements(self.interview_id)['requirements'][0]
            self.assertIsNone(fetched['specification'])





    def test_discovery_failure_does_not_erase_existing_requirements(self):
        req = self.save_requirement()
        original_id = req.id
        agent = RequirementsAgent()
        with patch.object(agent.ollama, 'generate_json', side_effect=RuntimeError('offline')):
            with self.assertRaises(RuntimeError):
                agent.process(self.interview_id)
        self.assertEqual(Requirement.query.one().id, original_id)

    def test_agent_persists_validated_details(self):
        agent = RequirementsAgent()
        candidate = {'description': 'El sistema deberá registrar libros con título y autor.',
                     'actor': 'bibliotecario', 'ambiguity_detected': False, 'verifiable': True}
        with patch.object(agent, '_extract', return_value=[candidate]), patch.object(
                agent.ollama, 'generate_json', return_value={'inputs': [1],
                'business_rules': [2, 500]}):
            result = agent.process(self.interview_id)
        self.assertEqual(len(result), 1)
        req = Requirement.query.one()
        fields = json.loads(req.specification)['fields']
        self.assertEqual(fields['inputs'], ['El bibliotecario debe registrar libros con título y autor.', 'El título es obligatorio.'])
        self.assertEqual(fields['business_rules'], ['El título es obligatorio.'])

    def test_complete_interview_is_one_function_with_all_components(self):
        source = (Path(__file__).parent / 'fixtures/library_registration.txt').read_text(encoding='utf-8').strip()
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription = source
        db.session.commit()
        agent = RequirementsAgent()
        # Adversarial model: assigns every sentence to every field. The
        # classifier must still keep results, conditions and test cases apart.
        from app.services.requirement_specification import FIELDS
        noisy = {key: list(range(1, 9)) for key in FIELDS}
        with patch.object(agent.ollama, 'generate_json', return_value=noisy), patch.object(agent, '_extract') as extract:
            result = agent.process(self.interview_id)
        extract.assert_not_called()
        self.assertEqual(len(result), 1)
        self.assertEqual(result[0]['actor'], 'bibliotecario')
        self.assertEqual(result[0]['priority'], 'esencial')
        self.assertIn('código de inventario', result[0]['description'])
        self.assertIn('año de publicación', result[0]['description'])
        req = Requirement.query.one()
        rows = {r['key']: r['values'] for r in req.specification_rows}
        self.assertTrue(all(rows.values()), rows)
        self.assertEqual(len(rows['acceptance_criteria']), 5)
        self.assertIn('no debe guardarse el libro', ' '.join(rows['acceptance_criteria']))
        self.assertIn('sesión', ' '.join(rows['preconditions']))
        self.assertIn('150', ' '.join(rows['inputs']))
        self.assertIn('único', ' '.join(rows['business_rules']))
        self.assertIn('cancela', ' '.join(rows['alternatives']))
        self.assertNotIn('cuenta activa', ' '.join(rows['inputs']))
        self.assertNotIn('queda almacenado', ' '.join(rows['preconditions']))
        html = self.app.test_client().get(f'/interviews/{self.interview_id}/requirements').get_data(as_text=True)
        self.assertNotIn('Por aclarar', html)

    def test_separate_functions_and_negative_obligations(self):
        source = ('El sistema deberá registrar libros con título y autor. El título es obligatorio. '
                  'El sistema deberá consultar préstamos por fecha. Los datos de entrada son dos fechas. '
                  'El sistema no debe permitir borrar préstamos activos.')
        from app.services.functional_context import functional_blocks
        # Negative direct obligations are primary requests too.
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription = source
        db.session.commit()
        agent = RequirementsAgent()
        with patch.object(agent.ollama, 'generate_json', return_value={}):
            result = agent.process(self.interview_id)
        self.assertEqual(len(result), 3)
        self.assertIn('no debe permitir', result[2]['description'])
        contexts = functional_blocks(source)
        self.assertNotIn('fechas', ' '.join(contexts[0]['sentences']))
        self.assertNotIn('título', ' '.join(contexts[1]['sentences']))

    def test_clinic_compound_and_independent_functions(self):
        source = ('El sistema deberá registrar pacientes, agendar citas y emitir facturas. '
                  'El sistema deberá consultar pacientes. '
                  'El sistema deberá registrar pagos de clientes. '
                  'El sistema deberá registrar pagos de proveedores.')
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription = source
        db.session.commit()
        agent = RequirementsAgent()
        with patch.object(agent.ollama, 'generate_json', return_value={}):
            result = agent.process(self.interview_id)
        self.assertEqual(len(result), 6)
        self.assertEqual(Requirement.query.count(), 6)
        self.assertEqual(result[1]['description'], 'El sistema deberá agendar citas.')
        for req in Requirement.query.all():
            text = ' '.join(q for row in req.specification_rows for q in row['values'])
            if 'facturas' in req.description:
                self.assertNotIn('pacientes', text)
                self.assertIn('facturas', text)

    def test_semantic_discovery_for_an_unrecognized_expression(self):
        source = 'Nos hace falta conservar un historial de tratamientos.'
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription = source
        db.session.commit()
        agent = RequirementsAgent()
        def response(prompt, schema):
            return {'starts': [0]} if 'starts' in schema['properties'] else {}
        with patch.object(agent.ollama, 'generate_json', side_effect=response), patch.object(
                agent, '_extract', return_value=[{'description': 'El sistema deberá conservar un historial de tratamientos.'}]):
            result = agent.process(self.interview_id)
        self.assertEqual(len(result), 1)
        self.assertEqual(json.loads(Requirement.query.one().specification)['scope'], [source])

    def test_four_functions_real_transcription_complete_and_pending(self):
        source = (Path(__file__).parent / 'fixtures/library_four_functions.txt').read_text(encoding='utf-8')
        from app.services.functional_context import functional_blocks
        self.assertEqual(len(functional_blocks(source)), 4)
        interview = db.session.get(Interview, self.interview_id)
        interview.transcription = source
        db.session.commit()
        agent = RequirementsAgent()
        def noisy_response(prompt, schema):
            if 'starts' in schema['properties']:
                return {'starts': schema['properties']['starts']['items']['enum']}
            return {key: list(range(1, 9)) for key in schema['properties']}
        with patch.object(agent.ollama, 'generate_json', side_effect=noisy_response):
            result = agent.process(self.interview_id)
        self.assertEqual(len(result), 4)
        self.assertEqual([r['priority'] for r in result], ['esencial', 'esencial', 'esencial', 'opcional'])
        saved = Requirement.query.order_by(Requirement.id).all()
        rows = [{r['key']: r['values'] for r in req.specification_rows} for req in saved]
        for fiche in rows[:2]:
            self.assertTrue(all(fiche.values()), fiche)
            self.assertEqual(len(fiche['acceptance_criteria']), 3)
        for fiche in rows[2:]:
            for field in ('inputs', 'preconditions', 'alternatives', 'acceptance_criteria'):
                self.assertFalse(fiche[field], (field, fiche[field]))
            self.assertTrue(fiche['outputs'])
        self.assertTrue(all(not r['verifiable'] and r['ambiguity_detected'] for r in result[2:]))
        self.assertIn('ordenada', ' '.join(rows[3]['business_rules']))
        all_quotes = [q for fiche in rows for values in fiche.values() for q in values]
        self.assertFalse(any('reunión adicional' in q for q in all_quotes))




if __name__ == '__main__':
    unittest.main()
