import re

from io import BytesIO

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH


class DocumentService:

    def create_docx(self, content):

        document = Document()

        # Página
        section = document.sections[0]

        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

        # Fuente normal
        normal = document.styles["Normal"]

        normal.font.name = "Times New Roman"
        normal.font.size = Pt(12)

        # Encabezados
        for style_name, size in (
            ("Heading 1", 14),
            ("Heading 2", 13),
            ("Heading 3", 12)
        ):

            style = document.styles[
                style_name
            ]

            style.font.name = "Times New Roman"
            style.font.size = Pt(size)
            style.font.bold = True

        # =================================================
        # CONTENIDO
        # =================================================

        first_title = True

        for raw_line in content.splitlines():

            line = raw_line.strip()

            if not line:
                continue

            if line == "---":
                continue

            # Título principal
            if (
                first_title
                and line.startswith("# ")
            ):

                paragraph = (
                    document.add_paragraph()
                )

                paragraph.alignment = (
                    WD_ALIGN_PARAGRAPH.CENTER
                )

                run = paragraph.add_run(
                    line[2:].strip()
                )

                run.bold = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(16)

                first_title = False

                continue

            # H1
            if line.startswith("# "):

                document.add_heading(
                    line[2:].strip(),
                    level=1
                )

                continue

            # H2
            if line.startswith("## "):

                document.add_heading(
                    line[3:].strip(),
                    level=2
                )

                continue

            # H3
            if line.startswith("### "):

                document.add_heading(
                    line[4:].strip(),
                    level=3
                )

                continue

            # Evidencia
            if line.startswith("> "):

                paragraph = (
                    document.add_paragraph()
                )

                paragraph.paragraph_format.left_indent = (
                    Inches(0.3)
                )

                run = paragraph.add_run(
                    line[2:].strip()
                )

                run.italic = True
                run.font.name = "Times New Roman"
                run.font.size = Pt(11)

                continue

            # Listas
            if line.startswith("- "):

                paragraph = document.add_paragraph(
                    line[2:].strip(),
                    style="List Bullet"
                )

                self._format_runs(
                    paragraph
                )

                continue

            # Actor, prioridad, etc.
            bold = re.match(
                r"\*\*(.+?)\*\*\s*(.*)",
                line
            )

            if bold:

                paragraph = (
                    document.add_paragraph()
                )

                label = paragraph.add_run(
                    bold.group(1)
                )

                label.bold = True
                label.font.name = "Times New Roman"
                label.font.size = Pt(12)

                if bold.group(2):

                    value = paragraph.add_run(
                        " " + bold.group(2)
                    )

                    value.font.name = (
                        "Times New Roman"
                    )

                    value.font.size = Pt(12)

                continue

            # Texto normal
            paragraph = document.add_paragraph(
                line
            )

            paragraph.alignment = (
                WD_ALIGN_PARAGRAPH.JUSTIFY
            )

            self._format_runs(
                paragraph
            )

        buffer = BytesIO()

        document.save(
            buffer
        )

        buffer.seek(0)

        return buffer

    def _format_runs(
        self,
        paragraph
    ):

        for run in paragraph.runs:

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)